#!/usr/bin/env python3
"""Fill the official FAIR Competence Funding template from proposal.qmd.

Writes export/2026-09-15-fair-by-doing-proposal.docx. Keeps the template's
own styles: text is written into the existing (text) placeholder paragraphs
and table cells, and the guiding-question blocks are deleted.

Needs python-docx. Run with a venv interpreter if the system Python is
externally managed:
    <venv>/bin/python tools/fill_template.py
"""

import csv
import re
from pathlib import Path

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parent.parent
QMD = ROOT / "proposal" / "proposal.qmd"
TMPL = ROOT / "references" / "FAIR-Seed-Funding-2ndcall-Template.docx"
TABLES = ROOT / "data" / "tables"
OUT = ROOT / "export" / "2026-09-15-fair-by-doing-proposal.docx"

BUDGET = [
    ("Applicant training and education", "0"),
    ("Travel", "0"),
    ("Equipment", "9,000"),
    ("Publications", "0"),
    ("Conferences and Workshop organisation", "40,880"),
    ("Other", "0"),
]
TOTALS = ["49,880", "0", "49,880"]


def strip_todo(text):
    """Remove [TODO ...] markers and pandoc escapes.

    The visual editor writes \\[, \\] and \\@ into the qmd. Quarto renders
    them, but text pasted straight into the DOCX must not carry the
    backslashes, so unescape first and strip the TODO markers second.
    """
    text = re.sub(r"\\([\[\]@])", r"\1", text)
    return re.sub(r"\[TODO[^\]]*\]\s*", "", text).strip()


def parse():
    """Pull the content out of the qmd, keyed by template slot."""
    src = QMD.read_text(encoding="utf-8")
    src = re.sub(r"<!--.*?-->", "", src, flags=re.S)  # drop the author comment
    body = src[src.index("## Proposal submission"):]

    out = {}

    # Header block bullets.
    hdr = body[:body.index("### Project summary")]
    for line in hdr.split("\n"):
        if not line.startswith("- "):
            continue
        key, _, val = line[2:].partition(": ")
        out.setdefault("header", {})[key.strip()] = val.strip()

    # Named blocks, split on headings.
    blocks = re.split(r"\n(?=#{2,4} )", body)
    applicant = None
    for blk in blocks:
        head = blk.split("\n")[0].strip()
        rest = strip_todo(blk[len(blk.split("\n")[0]):])
        if head.startswith("### Project summary"):
            out["summary"] = rest
        elif head.startswith("### Applicant 1"):
            applicant = "L"
        elif head.startswith("### Applicant 2"):
            applicant = "A"
        elif head.startswith("#### 1.1"):
            out[f"1.1{applicant}"] = rest
        elif head.startswith("#### 1.2"):
            out[f"1.2{applicant}"] = rest
        elif re.match(r"### 2\.[1-5] ", head):
            out[head.split()[1]] = rest
        elif head.startswith("### 3.1"):
            out["3.1"] = rest
        elif head.startswith("### 3.2"):
            out["3.2"] = rest
    return out


def clear(par):
    for r in list(par.runs):
        r._element.getparent().remove(r._element)


def write_para(par, text, bold=False, italic=False, size=10):
    """Replace a paragraph's content, keeping its style."""
    clear(par)
    run = par.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    return par


def para_after(par, text, style=None, bold=False, size=10):
    """Insert a new paragraph directly after `par` and return it."""
    new = par.insert_paragraph_before(text)
    # insert_paragraph_before puts it before, so move it after instead.
    par._element.addnext(new._element)
    clear(new)
    run = new.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if style:
        new.style = style
    return new


def fill_block(par, text):
    """Write a multi-paragraph block into a single (text) placeholder."""
    chunks = [c.strip() for c in text.split("\n\n") if c.strip()]
    if not chunks:
        return par
    lines = []
    for c in chunks:
        if c.lstrip().startswith("-"):
            lines.extend(l.strip().lstrip("- ").strip()
                         for l in c.split("\n") if l.strip())
        else:
            lines.append(c.replace("\n", " "))
    write_para(par, lines[0])
    cur = par
    for extra in lines[1:]:
        cur = para_after(cur, extra)
    return cur


def delete(par):
    par._element.getparent().remove(par._element)


def main():
    content = parse()
    doc = Document(str(TMPL))

    # --- header table -------------------------------------------------
    h = content["header"]
    t = doc.tables[0]
    rows = {
        0: h.get("Title of project", ""),
        1: content.get("summary", ""),
        2: h.get("Corresponding applicant", ""),
        3: h.get("Co-applicant", ""),
        4: h.get("Supporting research unit 1 (direct superior)", ""),
        5: h.get("Supporting research unit 2", ""),
    }
    for ri, val in rows.items():
        cell = t.rows[ri].cells[1]
        write_para(cell.paragraphs[0], val)

    # --- walk the body, filling placeholders and deleting prompts ------
    paras = doc.paragraphs
    # Map each (text) placeholder to its content key, in document order.
    slots = ["1.1L", "1.2L", "2.1", "2.2", "2.3", "2.4", "2.5", "3.2"]
    slot_i = 0
    to_delete = []
    pending_applicant2 = None

    # Guiding questions run from a "Guiding ...:" line up to the next
    # (text) placeholder. Matching on a trailing "?" misses the ones that
    # end in a parenthetical, so delete the whole span structurally.
    in_guiding = False
    for p in paras:
        txt = p.text.strip()
        if txt.lower().startswith("guiding") and txt.endswith(":"):
            in_guiding = True
            to_delete.append(p)
            continue
        if in_guiding:
            if txt == "(text)":
                in_guiding = False
            else:
                to_delete.append(p)
                continue
        if txt.startswith("Please use the format provided"):
            to_delete.append(p)
        elif txt == "(text)":
            key = slots[slot_i]
            slot_i += 1
            last = fill_block(p, content[key])
            # After Lars's 1.1 and 1.2, append Adriana's under her own
            # subheading, since the template carries the pair only once.
            if key == "1.1L":
                hd = para_after(last, "1.1 Adriana Clavijo Daza", bold=True)
                fill_block(para_after(hd, "x"), content["1.1A"])
            elif key == "1.2L":
                hd = para_after(last, "1.2 Adriana Clavijo Daza", bold=True)
                fill_block(para_after(hd, "x"), content["1.2A"])

    for p in to_delete:
        delete(p)

    # --- applicant name headings for Lars -----------------------------
    for p in doc.paragraphs:
        if p.text.strip().startswith("1.1 Work related to Research"):
            write_para(p, "1.1 Lars Schöbitz: work related to research, RDM, "
                          "RSE, ORD and the FAIR principles", bold=True)
        elif p.text.strip().startswith("1.2 Previous and ongoing"):
            write_para(p, "1.2 Lars Schöbitz: previous and ongoing RDM/RSE "
                          "activities at ETH Zurich", bold=True)

    # --- budget table --------------------------------------------------
    bt = doc.tables[1]
    for i, (_, amount) in enumerate(BUDGET, start=1):
        for ci in (2, 3):
            write_para(bt.rows[i].cells[ci].paragraphs[0], amount)
    for off, val in enumerate(TOTALS):
        for ci in (2, 3):
            write_para(bt.rows[7 + off].cells[ci].paragraphs[0], val,
                       bold=True)

    # --- work plan table into 2.3 --------------------------------------
    insert_work_plan(doc)

    OUT.parent.mkdir(exist_ok=True)
    doc.save(str(OUT))
    print(f"wrote {OUT}")
    return doc


def insert_work_plan(doc):
    """Append the milestone table at the end of the 2.3 block."""
    csv_path = TABLES / "tbl-01-work-packages.csv"
    if not csv_path.exists():
        return
    rows = list(csv.DictReader(csv_path.open(encoding="utf-8")))
    anchor = None
    seen_23 = False
    for p in doc.paragraphs:
        if p.text.strip().startswith("2.3 Work plan"):
            seen_23 = True
            continue
        if seen_23 and p.text.strip().startswith("2.4 "):
            anchor = p
            break
    if anchor is None:
        return

    # Months and Lead come from the sheet now; the legend for LS, AC and GW
    # sits on the roles sentence in 2.3.
    t = doc.add_table(rows=len(rows) + 1, cols=4)
    t.style = "Table Grid"
    hdr = ["WP", "Goal", "Months", "Lead"]
    for ci, htxt in enumerate(hdr):
        write_para(t.rows[0].cells[ci].paragraphs[0], htxt, bold=True, size=9)
    for ri, r in enumerate(rows, start=1):
        vals = [f'{r["WP"]}: {r["Name"]}', r["Goal"], r.get("Months", ""),
                r.get("Lead", "")]
        for ci, v in enumerate(vals):
            write_para(t.rows[ri].cells[ci].paragraphs[0], v, size=9)
    anchor._element.addprevious(t._element)


if __name__ == "__main__":
    main()
